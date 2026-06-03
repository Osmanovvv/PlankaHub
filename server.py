from __future__ import annotations

import hashlib
import hmac
import json
import sqlite3
import uuid
from io import BytesIO
from http.cookies import CookieError, SimpleCookie
from datetime import datetime, timezone
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse


ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
DB_PATH = DATA_DIR / "registrations.db"
ADMIN_USERNAME = "safksAJ"
ADMIN_PASSWORD = "SA_)-14AKsafdq_1"
ADMIN_COOKIE_NAME = "planka_admin"
ADMIN_PATH = "/asjkfhjkqwfasf14871209asjkSA"
ADMIN_LOGIN_PATH = f"{ADMIN_PATH}/login"
ADMIN_LOGOUT_PATH = f"{ADMIN_PATH}/logout"
COMMUNITY_BASE_PARTICIPANTS = 267
COMMUNITY_BASE_CARD_NUMBER = 268
COMMUNITY_COUNTER_BASELINE_KEY = "community_step2_baseline"


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def init_db() -> None:
    DATA_DIR.mkdir(exist_ok=True)
    with sqlite3.connect(DB_PATH) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS registration_attempts (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              attempt_uuid TEXT NOT NULL UNIQUE,
              status TEXT NOT NULL DEFAULT 'in_progress',
              current_step INTEGER NOT NULL DEFAULT 1,
              fields_json TEXT NOT NULL DEFAULT '{}',
              utm_json TEXT NOT NULL DEFAULT '{}',
              user_agent TEXT,
              referrer TEXT,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL,
              submitted_at TEXT
            )
            """
        )
        columns = {
            row[1]
            for row in connection.execute("PRAGMA table_info(registration_attempts)").fetchall()
        }
        if "utm_json" not in columns:
            connection.execute("ALTER TABLE registration_attempts ADD COLUMN utm_json TEXT NOT NULL DEFAULT '{}'")
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS site_settings (
              key TEXT PRIMARY KEY,
              value TEXT NOT NULL
            )
            """
        )
        baseline = connection.execute(
            "SELECT value FROM site_settings WHERE key = ?",
            (COMMUNITY_COUNTER_BASELINE_KEY,),
        ).fetchone()
        if baseline is None:
            current_step2_count = connection.execute(
                "SELECT COUNT(*) FROM registration_attempts WHERE current_step >= 2"
            ).fetchone()[0]
            connection.execute(
                "INSERT INTO site_settings (key, value) VALUES (?, ?)",
                (COMMUNITY_COUNTER_BASELINE_KEY, str(current_step2_count)),
            )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS behavior_sessions (
              session_id TEXT PRIMARY KEY,
              first_seen_at TEXT NOT NULL,
              last_seen_at TEXT NOT NULL,
              user_agent TEXT,
              referrer TEXT,
              landing_path TEXT,
              utm_json TEXT NOT NULL DEFAULT '{}',
              max_scroll_percent REAL NOT NULL DEFAULT 0,
              max_scroll_px INTEGER NOT NULL DEFAULT 0,
              viewport_width INTEGER,
              viewport_height INTEGER,
              page_height INTEGER,
              registration_attempt_id INTEGER
            )
            """
        )
        behavior_session_columns = {
            row[1]
            for row in connection.execute("PRAGMA table_info(behavior_sessions)").fetchall()
        }
        if "registration_attempt_id" not in behavior_session_columns:
            connection.execute("ALTER TABLE behavior_sessions ADD COLUMN registration_attempt_id INTEGER")
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS behavior_events (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              session_id TEXT NOT NULL,
              event_type TEXT NOT NULL,
              event_name TEXT NOT NULL,
              label TEXT,
              target_text TEXT,
              target_href TEXT,
              page_path TEXT,
              viewport_width INTEGER,
              viewport_height INTEGER,
              scroll_y INTEGER,
              scroll_percent REAL,
              metadata_json TEXT NOT NULL DEFAULT '{}',
              user_agent TEXT,
              referrer TEXT,
              created_at TEXT NOT NULL,
              registration_attempt_id INTEGER
            )
            """
        )
        behavior_event_columns = {
            row[1]
            for row in connection.execute("PRAGMA table_info(behavior_events)").fetchall()
        }
        if "registration_attempt_id" not in behavior_event_columns:
            connection.execute("ALTER TABLE behavior_events ADD COLUMN registration_attempt_id INTEGER")
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_behavior_events_created ON behavior_events(created_at)"
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_behavior_events_name ON behavior_events(event_type, event_name)"
        )
        connection.commit()


def row_to_dict(row: sqlite3.Row) -> dict:
    fields = {}
    if row["fields_json"]:
        try:
            fields = json.loads(row["fields_json"])
        except json.JSONDecodeError:
            fields = {}

    utm = {}
    if row["utm_json"]:
        try:
            utm = json.loads(row["utm_json"])
        except json.JSONDecodeError:
            utm = {}

    return {
        "id": row["id"],
        "displayId": f"#{row['id']}",
        "synthetic": False,
        "attemptUuid": row["attempt_uuid"],
        "status": row["status"],
        "currentStep": row["current_step"],
        "fields": fields,
        "utm": utm,
        "userAgent": row["user_agent"],
        "referrer": row["referrer"],
        "createdAt": row["created_at"],
        "updatedAt": row["updated_at"],
        "submittedAt": row["submitted_at"],
    }


def parse_json_object(value: str | None) -> dict:
    if not value:
        return {}

    try:
        parsed = json.loads(value)
    except json.JSONDecodeError:
        return {}

    return parsed if isinstance(parsed, dict) else {}


def session_row_to_attempt(row: sqlite3.Row) -> dict:
    session_id = row["session_id"]
    fields = {
        "landingPath": row["landing_path"],
        "maxScrollPercent": row["visible_scroll_percent"] or row["max_scroll_percent"] or 0,
        "clickCount": row["click_count"] or 0,
        "eventCount": row["event_count"] or 0,
    }

    return {
        "id": f"visit:{session_id}",
        "displayId": "визит",
        "synthetic": True,
        "attemptUuid": session_id,
        "status": "not_started",
        "currentStep": 0,
        "fields": fields,
        "utm": parse_json_object(row["utm_json"]),
        "userAgent": row["user_agent"],
        "referrer": row["referrer"],
        "createdAt": row["first_seen_at"],
        "updatedAt": row["last_seen_at"],
        "submittedAt": None,
    }


def load_admin_attempt_rows(connection: sqlite3.Connection) -> list[dict]:
    rows = connection.execute(
        "SELECT * FROM registration_attempts ORDER BY datetime(created_at) DESC, id DESC"
    ).fetchall()
    visit_rows = connection.execute(
        """
        WITH session_stats AS (
          SELECT
            s.*,
            MAX(e.registration_attempt_id) AS event_attempt_id,
            COUNT(CASE WHEN e.event_name != 'scroll.final' THEN e.id END) AS event_count,
            SUM(CASE WHEN e.event_type = 'click' THEN 1 ELSE 0 END) AS click_count,
            MAX(
              CASE
                WHEN e.event_type = 'scroll' AND e.event_name != 'scroll.final'
                THEN e.scroll_percent
              END
            ) AS visible_scroll_percent
          FROM behavior_sessions s
          LEFT JOIN behavior_events e ON e.session_id = s.session_id
          GROUP BY s.session_id
        )
        SELECT *
        FROM session_stats
        WHERE COALESCE(registration_attempt_id, event_attempt_id) IS NULL
        ORDER BY datetime(last_seen_at) DESC
        LIMIT 5000
        """
    ).fetchall()
    attempts = [row_to_dict(row) for row in rows]
    attempts.extend(session_row_to_attempt(row) for row in visit_rows)
    attempts.sort(key=lambda item: item.get("updatedAt") or item.get("createdAt") or "", reverse=True)
    return attempts


def parse_current_step(body: dict, default: int = 1) -> int:
    raw_step = body.get("currentStep")

    if raw_step is None:
        return default

    try:
        step = int(raw_step)
    except (TypeError, ValueError):
        return default

    return max(0, min(step, 3))


def attempt_status_label(item: dict) -> str:
    if item["status"] == "submitted":
        return "Отправлено"

    return "В процессе" if int(item["currentStep"] or 0) > 0 else "Не пройдена регистрация"


def attempt_step_label(item: dict) -> str:
    step = int(item["currentStep"] or 0)
    return f"{step} / 3" if step > 0 else "нет"


def safe_text(value, limit: int = 500) -> str | None:
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None

    return text[:limit]


def safe_int(value, default: int = 0, minimum: int = 0, maximum: int = 100000) -> int:
    try:
        number = int(float(value))
    except (TypeError, ValueError):
        return default

    return max(minimum, min(number, maximum))


def safe_float(value, default: float = 0.0, minimum: float = 0.0, maximum: float = 100.0) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return default

    return max(minimum, min(number, maximum))


def get_community_counter() -> dict:
    with sqlite3.connect(DB_PATH) as connection:
        current_step2_count = connection.execute(
            "SELECT COUNT(*) FROM registration_attempts WHERE current_step >= 2"
        ).fetchone()[0]
        baseline_row = connection.execute(
            "SELECT value FROM site_settings WHERE key = ?",
            (COMMUNITY_COUNTER_BASELINE_KEY,),
        ).fetchone()

        if baseline_row is None:
            baseline = current_step2_count
            connection.execute(
                "INSERT INTO site_settings (key, value) VALUES (?, ?)",
                (COMMUNITY_COUNTER_BASELINE_KEY, str(baseline)),
            )
            connection.commit()
        else:
            baseline = safe_int(baseline_row[0], default=0, maximum=1000000)

    step2_count = max(0, current_step2_count - baseline)
    participants = COMMUNITY_BASE_PARTICIPANTS + step2_count
    card_number = COMMUNITY_BASE_CARD_NUMBER + step2_count

    return {
        "baseParticipants": COMMUNITY_BASE_PARTICIPANTS,
        "baseCardNumber": COMMUNITY_BASE_CARD_NUMBER,
        "step2Count": step2_count,
        "participants": participants,
        "cardNumber": card_number,
    }


class PlankaHandler(SimpleHTTPRequestHandler):
    server_version = "PlankaHub/1.0"

    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def end_headers(self) -> None:
        self.send_header("Cache-Control", "no-store")
        super().end_headers()

    def do_GET(self) -> None:
        path = urlparse(self.path).path

        if path in {ADMIN_PATH, f"{ADMIN_PATH}/"}:
            self.path = "/admin.html" if self.is_admin_authenticated() else "/admin-login.html"
            super().do_GET()
            return

        if path in {"/admin", "/admin/"}:
            self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)
            return

        if path == ADMIN_LOGOUT_PATH:
            self.handle_admin_logout()
            return

        if self.is_protected_admin_path(path) and not self.require_admin_auth(path):
            return

        if path == "/api/admin/registration-attempts/export.xlsx":
            self.handle_export_attempts()
            return

        if path == "/api/admin/registration-attempts":
            self.handle_list_attempts()
            return

        if path == "/api/admin/behavior-analytics":
            self.handle_behavior_analytics()
            return

        if path == "/api/community-counter":
            self.handle_community_counter()
            return

        super().do_GET()

    def do_POST(self) -> None:
        path = urlparse(self.path).path

        if path == ADMIN_LOGIN_PATH:
            self.handle_admin_login()
            return

        if path == ADMIN_LOGOUT_PATH:
            self.handle_admin_logout()
            return

        if path == "/api/analytics/events":
            self.handle_create_behavior_event()
            return

        if self.is_protected_admin_path(path) and not self.require_admin_auth(path):
            return

        if path == "/api/registration-attempts":
            self.handle_create_attempt()
            return

        if path.startswith("/api/registration-attempts/") and path.endswith("/submit"):
            self.handle_submit_attempt(path)
            return

        self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def do_PATCH(self) -> None:
        path = urlparse(self.path).path

        if path.startswith("/api/registration-attempts/"):
            self.handle_update_attempt(path)
            return

        self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def do_DELETE(self) -> None:
        path = urlparse(self.path).path

        if self.is_protected_admin_path(path) and not self.require_admin_auth(path):
            return

        if path == "/api/admin/registration-attempts":
            self.handle_delete_attempts()
            return

        if path.startswith("/api/admin/registration-attempts/"):
            self.handle_delete_attempt(path)
            return

        self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def read_json(self) -> dict:
        length = int(self.headers.get("Content-Length", "0") or "0")
        if length == 0:
            return {}

        raw_body = self.rfile.read(length).decode("utf-8")
        try:
            parsed = json.loads(raw_body)
        except json.JSONDecodeError:
            return {}

        return parsed if isinstance(parsed, dict) else {}

    def read_form(self) -> dict[str, str]:
        length = int(self.headers.get("Content-Length", "0") or "0")
        if length == 0:
            return {}

        raw_body = self.rfile.read(length).decode("utf-8")
        parsed = parse_qs(raw_body, keep_blank_values=True)
        return {key: values[0] for key, values in parsed.items() if values}

    def send_json(self, payload: dict | list, status: HTTPStatus = HTTPStatus.OK) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    @staticmethod
    def is_protected_admin_path(path: str) -> bool:
        return path == "/admin.html" or path.startswith("/api/admin/")

    @staticmethod
    def admin_auth_token() -> str:
        return hmac.new(
            ADMIN_PASSWORD.encode("utf-8"),
            ADMIN_USERNAME.encode("utf-8"),
            hashlib.sha256,
        ).hexdigest()

    def is_admin_authenticated(self) -> bool:
        try:
            cookie = SimpleCookie(self.headers.get("Cookie", ""))
        except CookieError:
            return False

        auth_cookie = cookie.get(ADMIN_COOKIE_NAME)
        if auth_cookie is None:
            return False

        return hmac.compare_digest(auth_cookie.value, self.admin_auth_token())

    def require_admin_auth(self, path: str) -> bool:
        if self.is_admin_authenticated():
            return True

        if path.startswith("/api/admin/"):
            self.send_json({"error": "Unauthorized"}, HTTPStatus.UNAUTHORIZED)
        else:
            self.redirect(ADMIN_PATH)

        return False

    def redirect(self, location: str, status: HTTPStatus = HTTPStatus.SEE_OTHER) -> None:
        self.send_response(status)
        self.send_header("Location", location)
        self.send_header("Content-Length", "0")
        self.end_headers()

    def handle_admin_login(self) -> None:
        form = self.read_form()
        username = form.get("username", "")
        password = form.get("password", "")

        if (
            hmac.compare_digest(username, ADMIN_USERNAME)
            and hmac.compare_digest(password, ADMIN_PASSWORD)
        ):
            self.send_response(HTTPStatus.SEE_OTHER)
            self.send_header("Location", ADMIN_PATH)
            self.send_header(
                "Set-Cookie",
                f"{ADMIN_COOKIE_NAME}={self.admin_auth_token()}; Path=/; HttpOnly; SameSite=Lax",
            )
            self.send_header("Content-Length", "0")
            self.end_headers()
            return

        self.redirect(f"{ADMIN_PATH}?error=1")

    def handle_admin_logout(self) -> None:
        self.send_response(HTTPStatus.SEE_OTHER)
        self.send_header("Location", ADMIN_PATH)
        self.send_header(
            "Set-Cookie",
            f"{ADMIN_COOKIE_NAME}=; Path=/; Max-Age=0; HttpOnly; SameSite=Lax",
        )
        self.send_header("Content-Length", "0")
        self.end_headers()

    def handle_create_attempt(self) -> None:
        body = self.read_json()
        now = utc_now()
        attempt_uuid = str(uuid.uuid4())
        current_step = parse_current_step(body)
        fields = body.get("fields") if isinstance(body.get("fields"), dict) else {}
        utm = body.get("utm") if isinstance(body.get("utm"), dict) else {}

        with sqlite3.connect(DB_PATH) as connection:
            connection.row_factory = sqlite3.Row
            cursor = connection.execute(
                """
                INSERT INTO registration_attempts
                  (attempt_uuid, status, current_step, fields_json, utm_json, user_agent, referrer, created_at, updated_at)
                VALUES (?, 'in_progress', ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    attempt_uuid,
                    current_step,
                    json.dumps(fields, ensure_ascii=False),
                    json.dumps(utm, ensure_ascii=False),
                    self.headers.get("User-Agent"),
                    self.headers.get("Referer"),
                    now,
                    now,
                ),
            )
            row = connection.execute(
                "SELECT * FROM registration_attempts WHERE id = ?",
                (cursor.lastrowid,),
            ).fetchone()

        payload = row_to_dict(row)
        if int(payload.get("currentStep") or 0) >= 2:
            payload["communityCounter"] = get_community_counter()

        self.send_json(payload, HTTPStatus.CREATED)

    def handle_update_attempt(self, path: str) -> None:
        attempt_id = self.extract_attempt_id(path)
        if attempt_id is None:
            self.send_json({"error": "Invalid attempt id"}, HTTPStatus.BAD_REQUEST)
            return

        body = self.read_json()
        now = utc_now()
        current_step = parse_current_step(body)
        fields = body.get("fields") if isinstance(body.get("fields"), dict) else {}
        status = body.get("status") if body.get("status") in {"in_progress", "submitted"} else "in_progress"

        with sqlite3.connect(DB_PATH) as connection:
            connection.row_factory = sqlite3.Row
            existing = connection.execute(
                "SELECT status, current_step FROM registration_attempts WHERE id = ?",
                (attempt_id,),
            ).fetchone()

            if existing is None:
                self.send_json({"error": "Attempt not found"}, HTTPStatus.NOT_FOUND)
                return

            next_status = "submitted" if existing["status"] == "submitted" else status
            next_step = max(int(existing["current_step"] or 0), current_step)

            connection.execute(
                """
                UPDATE registration_attempts
                SET current_step = ?,
                    fields_json = ?,
                    status = ?,
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    next_step,
                    json.dumps(fields, ensure_ascii=False),
                    next_status,
                    now,
                    attempt_id,
                ),
            )
            row = connection.execute(
                "SELECT * FROM registration_attempts WHERE id = ?",
                (attempt_id,),
            ).fetchone()

        if row is None:
            self.send_json({"error": "Attempt not found"}, HTTPStatus.NOT_FOUND)
            return

        payload = row_to_dict(row)
        if int(payload.get("currentStep") or 0) >= 2:
            payload["communityCounter"] = get_community_counter()

        self.send_json(payload)

    def handle_submit_attempt(self, path: str) -> None:
        attempt_id = self.extract_attempt_id(path.removesuffix("/submit"))
        if attempt_id is None:
            self.send_json({"error": "Invalid attempt id"}, HTTPStatus.BAD_REQUEST)
            return

        body = self.read_json()
        now = utc_now()
        fields = body.get("fields") if isinstance(body.get("fields"), dict) else {}

        with sqlite3.connect(DB_PATH) as connection:
            connection.row_factory = sqlite3.Row
            connection.execute(
                """
                UPDATE registration_attempts
                SET current_step = 3,
                    fields_json = ?,
                    status = 'submitted',
                    updated_at = ?,
                    submitted_at = ?
                WHERE id = ?
                """,
                (
                    json.dumps(fields, ensure_ascii=False),
                    now,
                    now,
                    attempt_id,
                ),
            )
            row = connection.execute(
                "SELECT * FROM registration_attempts WHERE id = ?",
                (attempt_id,),
            ).fetchone()

        if row is None:
            self.send_json({"error": "Attempt not found"}, HTTPStatus.NOT_FOUND)
            return

        payload = row_to_dict(row)
        payload["communityCounter"] = get_community_counter()
        self.send_json(payload)

    def handle_community_counter(self) -> None:
        self.send_json(get_community_counter())

    def handle_create_behavior_event(self) -> None:
        body = self.read_json()
        session_id = safe_text(body.get("sessionId"), 80)
        event_type = safe_text(body.get("type"), 40) or "click"
        event_name = safe_text(body.get("name"), 120)

        if not session_id or not event_name:
            self.send_json({"error": "sessionId and name are required"}, HTTPStatus.BAD_REQUEST)
            return

        if event_type not in {"click", "scroll", "registration", "visit"}:
            event_type = "click"

        now = utc_now()
        metadata = body.get("metadata") if isinstance(body.get("metadata"), dict) else {}
        utm = body.get("utm") if isinstance(body.get("utm"), dict) else {}
        viewport = body.get("viewport") if isinstance(body.get("viewport"), dict) else {}
        scroll = body.get("scroll") if isinstance(body.get("scroll"), dict) else {}

        viewport_width = safe_int(viewport.get("width"), maximum=10000)
        viewport_height = safe_int(viewport.get("height"), maximum=10000)
        page_height = safe_int(scroll.get("pageHeight"), maximum=1000000)
        scroll_y = safe_int(scroll.get("y"), maximum=1000000)
        scroll_percent = safe_float(scroll.get("percent"))
        max_scroll_px = safe_int(scroll.get("maxY"), scroll_y, maximum=1000000)
        landing_path = safe_text(body.get("landingPath"), 500) or safe_text(body.get("pagePath"), 500) or "/"
        registration_attempt_id = safe_int(
            body.get("registrationAttemptId") or metadata.get("attemptId"),
            default=0,
            minimum=0,
            maximum=1000000000,
        ) or None

        with sqlite3.connect(DB_PATH) as connection:
            connection.execute(
                """
                INSERT INTO behavior_sessions
                  (session_id, first_seen_at, last_seen_at, user_agent, referrer, landing_path, utm_json,
                   max_scroll_percent, max_scroll_px, viewport_width, viewport_height, page_height,
                   registration_attempt_id)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(session_id) DO UPDATE SET
                  last_seen_at = excluded.last_seen_at,
                  referrer = COALESCE(behavior_sessions.referrer, excluded.referrer),
                  landing_path = COALESCE(behavior_sessions.landing_path, excluded.landing_path),
                  utm_json = CASE
                    WHEN excluded.utm_json IS NOT NULL AND excluded.utm_json != '{}'
                    THEN excluded.utm_json
                    ELSE behavior_sessions.utm_json
                  END,
                  max_scroll_percent = MAX(behavior_sessions.max_scroll_percent, excluded.max_scroll_percent),
                  max_scroll_px = MAX(behavior_sessions.max_scroll_px, excluded.max_scroll_px),
                  viewport_width = COALESCE(excluded.viewport_width, behavior_sessions.viewport_width),
                  viewport_height = COALESCE(excluded.viewport_height, behavior_sessions.viewport_height),
                  page_height = MAX(COALESCE(behavior_sessions.page_height, 0), COALESCE(excluded.page_height, 0)),
                  registration_attempt_id = COALESCE(
                    excluded.registration_attempt_id,
                    behavior_sessions.registration_attempt_id
                  )
                """,
                (
                    session_id,
                    now,
                    now,
                    self.headers.get("User-Agent"),
                    self.headers.get("Referer"),
                    landing_path,
                    json.dumps(utm, ensure_ascii=False),
                    scroll_percent,
                    max_scroll_px,
                    viewport_width or None,
                    viewport_height or None,
                    page_height or None,
                    registration_attempt_id,
                ),
            )
            connection.execute(
                """
                INSERT INTO behavior_events
                  (session_id, event_type, event_name, label, target_text, target_href, page_path,
                   viewport_width, viewport_height, scroll_y, scroll_percent, metadata_json,
                   user_agent, referrer, created_at, registration_attempt_id)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    session_id,
                    event_type,
                    event_name,
                    safe_text(body.get("label"), 180),
                    safe_text(body.get("text"), 500),
                    safe_text(body.get("href"), 500),
                    safe_text(body.get("pagePath"), 500),
                    viewport_width or None,
                    viewport_height or None,
                    scroll_y,
                    scroll_percent,
                    json.dumps(metadata, ensure_ascii=False),
                    self.headers.get("User-Agent"),
                    self.headers.get("Referer"),
                    now,
                    registration_attempt_id,
                ),
            )
            connection.commit()

        self.send_json({"ok": True}, HTTPStatus.CREATED)

    def handle_list_attempts(self) -> None:
        with sqlite3.connect(DB_PATH) as connection:
            connection.row_factory = sqlite3.Row
            attempts = load_admin_attempt_rows(connection)
        stats = {
            "total": len(attempts),
            "step0": sum(1 for item in attempts if item["currentStep"] == 0 and item["status"] != "submitted"),
            "step1": sum(1 for item in attempts if item["currentStep"] == 1 and item["status"] != "submitted"),
            "step2": sum(1 for item in attempts if item["currentStep"] == 2 and item["status"] != "submitted"),
            "step3": sum(1 for item in attempts if item["currentStep"] == 3 and item["status"] != "submitted"),
            "submitted": sum(1 for item in attempts if item["status"] == "submitted"),
        }

        self.send_json({"attempts": attempts, "stats": stats})

    def handle_behavior_analytics(self) -> None:
        with sqlite3.connect(DB_PATH) as connection:
            connection.row_factory = sqlite3.Row
            sessions = connection.execute(
                """
                WITH session_stats AS (
                  SELECT
                    s.*,
                    MAX(e.registration_attempt_id) AS event_attempt_id,
                    COUNT(CASE WHEN e.event_name != 'scroll.final' THEN e.id END) AS event_count,
                    SUM(CASE WHEN e.event_type = 'click' THEN 1 ELSE 0 END) AS click_count,
                    MAX(
                      CASE
                        WHEN e.event_type = 'scroll' AND e.event_name != 'scroll.final'
                        THEN e.scroll_percent
                      END
                    ) AS visible_scroll_percent
                  FROM behavior_sessions s
                  LEFT JOIN behavior_events e ON e.session_id = s.session_id
                  GROUP BY s.session_id
                )
                SELECT
                  session_stats.*,
                  a.id AS attempt_id,
                  a.attempt_uuid AS attempt_uuid,
                  a.status AS attempt_status,
                  a.current_step AS attempt_step,
                  a.submitted_at AS attempt_submitted_at
                FROM session_stats
                LEFT JOIN registration_attempts a
                  ON a.id = COALESCE(session_stats.registration_attempt_id, session_stats.event_attempt_id)
                ORDER BY datetime(session_stats.last_seen_at) DESC
                LIMIT 200
                """
            ).fetchall()
            top_clicks = connection.execute(
                """
                SELECT event_name, COALESCE(label, event_name) AS label, COUNT(*) AS count, MAX(created_at) AS last_at
                FROM behavior_events
                WHERE event_type = 'click'
                GROUP BY event_name, COALESCE(label, event_name)
                ORDER BY count DESC, datetime(last_at) DESC
                LIMIT 50
                """
            ).fetchall()
            recent_events = connection.execute(
                """
                SELECT
                  e.id,
                  e.session_id,
                  e.event_type,
                  e.event_name,
                  e.label,
                  e.target_text,
                  e.target_href,
                  e.page_path,
                  e.viewport_width,
                  e.viewport_height,
                  e.scroll_y,
                  e.scroll_percent,
                  e.created_at,
                  s.utm_json,
                  s.referrer AS session_referrer,
                  s.landing_path,
                  COALESCE(e.registration_attempt_id, s.registration_attempt_id) AS linked_attempt_id,
                  a.attempt_uuid AS attempt_uuid,
                  a.status AS attempt_status,
                  a.current_step AS attempt_step
                FROM behavior_events e
                LEFT JOIN behavior_sessions s ON s.session_id = e.session_id
                LEFT JOIN registration_attempts a
                  ON a.id = COALESCE(e.registration_attempt_id, s.registration_attempt_id)
                WHERE e.event_name != 'scroll.final'
                ORDER BY datetime(e.created_at) DESC, e.id DESC
                LIMIT 100
                """
            ).fetchall()
            session_event_rows = []
            if sessions:
                session_ids = [row["session_id"] for row in sessions]
                placeholders = ",".join("?" for _ in session_ids)
                session_event_rows = connection.execute(
                    f"""
                    SELECT
                      session_id,
                      event_type,
                      event_name,
                      COALESCE(label, event_name) AS label,
                      target_text,
                      target_href,
                      scroll_percent,
                      created_at
                    FROM behavior_events
                    WHERE session_id IN ({placeholders})
                      AND event_name != 'scroll.final'
                    ORDER BY datetime(created_at) ASC, id ASC
                    """,
                    session_ids,
                ).fetchall()

        events_by_session = {}
        for event_row in session_event_rows:
            events_by_session.setdefault(event_row["session_id"], []).append(
                {
                    "type": event_row["event_type"],
                    "name": event_row["event_name"],
                    "label": event_row["label"],
                    "text": event_row["target_text"],
                    "href": event_row["target_href"],
                    "scrollPercent": event_row["scroll_percent"],
                    "createdAt": event_row["created_at"],
                }
            )

        session_items = []
        for row in sessions:
            utm = parse_json_object(row["utm_json"])
            source = self.normalized_source({"utm": utm, "referrer": row["referrer"]})
            attempt_id = row["attempt_id"] or row["registration_attempt_id"] or row["event_attempt_id"]
            attempt_status = row["attempt_status"] or ("in_progress" if attempt_id else "not_started")

            session_items.append(
                {
                    "sessionId": row["session_id"],
                    "displayUuid": row["attempt_uuid"] or row["session_id"],
                    "firstSeenAt": row["first_seen_at"],
                    "lastSeenAt": row["last_seen_at"],
                    "referrer": row["referrer"],
                    "landingPath": row["landing_path"],
                    "utm": utm,
                    "source": source,
                    "maxScrollPercent": row["visible_scroll_percent"] or 0,
                    "maxScrollPx": row["max_scroll_px"],
                    "viewportWidth": row["viewport_width"],
                    "viewportHeight": row["viewport_height"],
                    "pageHeight": row["page_height"],
                    "eventCount": row["event_count"] or 0,
                    "clickCount": row["click_count"] or 0,
                    "registrationAttemptId": attempt_id,
                    "registrationAttemptUuid": row["attempt_uuid"],
                    "registrationStatus": attempt_status,
                    "registrationStep": row["attempt_step"] or 0,
                    "registrationCompleted": attempt_status == "submitted",
                    "registrationSubmittedAt": row["attempt_submitted_at"],
                    "events": events_by_session.get(row["session_id"], [])[-40:],
                }
            )

        total_sessions = len(session_items)
        total_clicks = sum(item["clickCount"] for item in session_items)
        sessions_with_utm = sum(1 for item in session_items if item["utm"])
        submitted_sessions = sum(1 for item in session_items if item["registrationCompleted"])
        average_scroll = (
            sum(float(item["maxScrollPercent"] or 0) for item in session_items) / total_sessions
            if total_sessions
            else 0
        )
        scroll_buckets = {
            "25": sum(1 for item in session_items if float(item["maxScrollPercent"] or 0) >= 25),
            "50": sum(1 for item in session_items if float(item["maxScrollPercent"] or 0) >= 50),
            "75": sum(1 for item in session_items if float(item["maxScrollPercent"] or 0) >= 75),
            "90": sum(1 for item in session_items if float(item["maxScrollPercent"] or 0) >= 90),
            "100": sum(1 for item in session_items if float(item["maxScrollPercent"] or 0) >= 99),
        }
        source_rows = []
        for source in sorted({item["source"] for item in session_items}):
            source_sessions = [item for item in session_items if item["source"] == source]
            source_rows.append(
                {
                    "source": source,
                    "sessions": len(source_sessions),
                    "clicks": sum(item["clickCount"] for item in source_sessions),
                    "submitted": sum(1 for item in source_sessions if item["registrationCompleted"]),
                    "averageScroll": round(
                        sum(float(item["maxScrollPercent"] or 0) for item in source_sessions)
                        / len(source_sessions),
                        1,
                    ),
                }
            )
        event_items = []
        for row in recent_events:
            utm = parse_json_object(row["utm_json"])
            source = self.normalized_source({"utm": utm, "referrer": row["session_referrer"]})
            attempt_id = row["linked_attempt_id"]
            attempt_status = row["attempt_status"] or ("in_progress" if attempt_id else "not_started")
            item = dict(row)
            item.update(
                {
                    "utm": utm,
                    "source": source,
                    "registrationAttemptId": attempt_id,
                    "registrationAttemptUuid": row["attempt_uuid"],
                    "registrationStatus": attempt_status,
                    "registrationStep": row["attempt_step"] or 0,
                    "registrationCompleted": attempt_status == "submitted",
                }
            )
            event_items.append(item)

        self.send_json(
            {
                "stats": {
                    "sessions": total_sessions,
                    "clicks": total_clicks,
                    "sessionsWithUtm": sessions_with_utm,
                    "submittedSessions": submitted_sessions,
                    "averageScroll": round(average_scroll, 1),
                    "scrollBuckets": scroll_buckets,
                },
                "topClicks": [dict(row) for row in top_clicks],
                "sources": source_rows,
                "sessions": session_items,
                "recentEvents": event_items,
            }
        )

    def handle_export_attempts(self) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.chart import BarChart, PieChart, Reference
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            self.send_json(
                {"error": "openpyxl is required to export Excel files"},
                HTTPStatus.INTERNAL_SERVER_ERROR,
            )
            return

        with sqlite3.connect(DB_PATH) as connection:
            connection.row_factory = sqlite3.Row
            attempts = load_admin_attempt_rows(connection)
            behavior_sessions = connection.execute(
                """
                WITH session_stats AS (
                  SELECT
                    s.*,
                    MAX(e.registration_attempt_id) AS event_attempt_id,
                    COUNT(CASE WHEN e.event_name != 'scroll.final' THEN e.id END) AS event_count,
                    SUM(CASE WHEN e.event_type = 'click' THEN 1 ELSE 0 END) AS click_count,
                    MAX(
                      CASE
                        WHEN e.event_type = 'scroll' AND e.event_name != 'scroll.final'
                        THEN e.scroll_percent
                      END
                    ) AS visible_scroll_percent
                  FROM behavior_sessions s
                  LEFT JOIN behavior_events e ON e.session_id = s.session_id
                  GROUP BY s.session_id
                )
                SELECT
                  session_stats.*,
                  a.id AS attempt_id,
                  a.attempt_uuid AS attempt_uuid,
                  a.status AS attempt_status,
                  a.current_step AS attempt_step,
                  a.submitted_at AS attempt_submitted_at
                FROM session_stats
                LEFT JOIN registration_attempts a
                  ON a.id = COALESCE(session_stats.registration_attempt_id, session_stats.event_attempt_id)
                ORDER BY datetime(session_stats.last_seen_at) DESC
                LIMIT 5000
                """
            ).fetchall()
            behavior_events = connection.execute(
                """
                SELECT
                  e.*,
                  s.utm_json AS session_utm_json,
                  s.referrer AS session_referrer,
                  s.landing_path,
                  COALESCE(e.registration_attempt_id, s.registration_attempt_id) AS linked_attempt_id,
                  a.attempt_uuid AS attempt_uuid,
                  a.status AS attempt_status,
                  a.current_step AS attempt_step
                FROM behavior_events e
                LEFT JOIN behavior_sessions s ON s.session_id = e.session_id
                LEFT JOIN registration_attempts a
                  ON a.id = COALESCE(e.registration_attempt_id, s.registration_attempt_id)
                WHERE e.event_name != 'scroll.final'
                ORDER BY datetime(e.created_at) DESC, e.id DESC
                LIMIT 10000
                """
            ).fetchall()

        known_field_keys = [
            "first-name",
            "last-name",
            "role",
            "stage",
            "needs",
            "request",
            "email",
            "contact",
            "privacy",
        ]
        known_utm_keys = [
            "utm_source",
            "utm_medium",
            "utm_campaign",
            "utm_content",
            "utm_term",
            "utm_id",
            "gclid",
            "yclid",
            "fbclid",
            "utm_region",
        ]
        extra_field_keys = sorted({
            key
            for attempt in attempts
            for key in attempt["fields"].keys()
            if key not in known_field_keys
        })
        extra_utm_keys = sorted({
            key
            for attempt in attempts
            for key in attempt["utm"].keys()
            if key not in known_utm_keys
        })

        columns = [
            ("id", "ID", lambda item: item["id"]),
            ("uuid", "UUID", lambda item: item["attemptUuid"]),
            ("status", "Статус", attempt_status_label),
            ("step", "Шаг", attempt_step_label),
            ("created_at", "Создано", lambda item: item["createdAt"]),
            ("updated_at", "Обновлено", lambda item: item["updatedAt"]),
            ("submitted_at", "Отправлено", lambda item: item["submittedAt"]),
            ("referrer", "Referrer", lambda item: item["referrer"]),
            ("user_agent", "User Agent", lambda item: item["userAgent"]),
        ]
        columns.extend((key, key, lambda item, field_key=key: item["utm"].get(field_key)) for key in known_utm_keys)
        columns.extend((key, key, lambda item, utm_key=key: item["utm"].get(utm_key)) for key in extra_utm_keys)
        columns.extend((key, key, lambda item, field_key=key: item["fields"].get(field_key)) for key in known_field_keys)
        columns.extend((key, key, lambda item, field_key=key: item["fields"].get(field_key)) for key in extra_field_keys)
        columns.extend([
            ("utm_json", "Все UTM JSON", lambda item: item["utm"]),
            ("fields_json", "Все поля JSON", lambda item: item["fields"]),
            ("source_norm", "ИсточникНорм", self.normalized_source),
            ("step_num", "ШагNum", lambda item: item["currentStep"]),
        ])

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Попытки регистрации"
        sheet.append([title for _key, title, _getter in columns])

        for attempt in attempts:
            sheet.append([self.excel_value(getter(attempt)) for _key, _title, getter in columns])

        header_fill = PatternFill("solid", fgColor="4475F2")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 28

        for column_index, column_cells in enumerate(sheet.columns, start=1):
            max_length = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))

            width = min(max(max_length + 2, 10), 42)
            sheet.column_dimensions[get_column_letter(column_index)].width = width

        column_letters = {
            key: get_column_letter(index)
            for index, (key, _title, _getter) in enumerate(columns, start=1)
        }
        self.add_analytics_sheet_v2(
            workbook,
            attempts,
            column_letters,
            Alignment,
            BarChart,
            Border,
            Font,
            PatternFill,
            PieChart,
            Reference,
            Side,
        )
        self.add_behavior_excel_sheets(
            workbook,
            behavior_sessions,
            behavior_events,
            Alignment,
            Font,
            PatternFill,
            get_column_letter,
        )

        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True

        output = BytesIO()
        workbook.save(output)
        body = output.getvalue()
        filename = f"registration-attempts-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"

        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    @staticmethod
    def parse_json_object(value: str | None) -> dict:
        if not value:
            return {}

        try:
            parsed = json.loads(value)
        except json.JSONDecodeError:
            return {}

        return parsed if isinstance(parsed, dict) else {}

    @staticmethod
    def behavior_action_label(row: sqlite3.Row | dict) -> str:
        name = row["event_name"] if row["event_name"] else ""
        scroll_percent = row["scroll_percent"] if "scroll_percent" in row.keys() else None
        labels = {
            "header.burger": "Бургер меню",
            "burger.logo": "Бургер: логотип",
            "burger.nav": "Бургер: пункт меню",
            "burger.access": "Бургер: получить доступ",
            "burger.telegram": "Бургер: Telegram",
            "header.logo": "Логотип",
            "header.login": "Вход",
            "hero.free": "Получить бесплатно",
            "founders.join": "Стать участником",
            "bloggers.arrow": "Стрелка блогеров",
            "report.join": "Стать участником",
            "pablo.apply": "Подать заявку в Telegram",
            "pablo.site": "Сайт HIT Venture",
            "reviews.arrow": "Стрелка отзывов",
            "telegram.subscribe": "Подписаться на канал",
            "faq.question": "Вопрос FAQ",
            "footer.anchor": "Ссылка в футере",
            "registration.attempt_created": "Начал регистрацию",
            "visit.landing": "Визит по UTM",
        }

        if name.startswith("scroll."):
            if name == "scroll.final":
                return f"Финальный скролл {round(float(scroll_percent or 0))}%"

            return f"Скролл {name.removeprefix('scroll.')}%"

        if name.startswith("burger."):
            text = safe_text(row["target_text"] if "target_text" in row.keys() else None, 120)
            href = safe_text(row["target_href"] if "target_href" in row.keys() else None, 500)
            anchor_labels = {
                "#ecosystem": "Экосистема",
                "#education": "Обучение",
                "#investor": "Стать инвестором",
                "#about": "О нас",
                "#faq": "Вопросы",
                "#telegram": "Телеграм",
            }
            if not text or "?" in text:
                for anchor, title in anchor_labels.items():
                    if href.endswith(anchor) or anchor in href:
                        text = title
                        break
            if name == "burger.logo":
                return f"Бургер: {text or 'логотип'}"
            if name == "burger.access":
                return f"Бургер: {text or 'получить доступ'}"
            if name == "burger.telegram":
                return f"Бургер: {text or 'Telegram'}"
            return f"Бургер: {text or 'пункт меню'}"

        return labels.get(name, row["label"] or name)

    @staticmethod
    def registration_label(status: str | None, attempt_id=None, step=0) -> str:
        if status == "submitted":
            return "Заявка отправлена"

        if attempt_id:
            return f"Не завершил, шаг {int(step or 0)}/3"

        return "Не начинал"

    def add_behavior_excel_sheets(
        self,
        workbook,
        behavior_sessions: list[sqlite3.Row],
        behavior_events: list[sqlite3.Row],
        Alignment,
        Font,
        PatternFill,
        get_column_letter,
    ) -> None:
        header_fill = PatternFill("solid", fgColor="4475F2")
        header_font = Font(color="FFFFFF", bold=True)

        def setup_sheet(sheet) -> None:
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 28

            for column_index, column_cells in enumerate(sheet.columns, start=1):
                max_length = 0
                for cell in column_cells:
                    value = "" if cell.value is None else str(cell.value)
                    max_length = max(max_length, len(value))

                width = min(max(max_length + 2, 10), 46)
                sheet.column_dimensions[get_column_letter(column_index)].width = width

        events_by_session: dict[str, list[sqlite3.Row]] = {}
        for event in sorted(behavior_events, key=lambda row: row["created_at"] or ""):
            events_by_session.setdefault(event["session_id"], []).append(event)

        sessions_sheet = workbook.create_sheet("Поведение сессии")
        sessions_sheet.append([
            "UUID заявки",
            "Сессия",
            "Источник",
            "UTM source",
            "UTM medium",
            "UTM campaign",
            "UTM content",
            "UTM term",
            "Referrer",
            "Landing path",
            "Первый визит",
            "Последний визит",
            "Прокрутка %",
            "Клики",
            "События",
            "Что нажимал",
            "Пороги скролла",
            "Регистрация",
            "ID заявки",
            "Шаг заявки",
            "Viewport",
            "Page height",
            "User Agent",
        ])

        for row in behavior_sessions:
            utm = self.parse_json_object(row["utm_json"])
            source = self.normalized_source({"utm": utm, "referrer": row["referrer"]})
            attempt_id = row["attempt_id"] or row["registration_attempt_id"] or row["event_attempt_id"]
            attempt_status = row["attempt_status"] or ("in_progress" if attempt_id else "not_started")
            attempt_step = row["attempt_step"] or 0
            session_events = events_by_session.get(row["session_id"], [])
            clicked = []
            scroll_milestones = []

            for event in session_events:
                action = self.behavior_action_label(event)
                if event["event_type"] == "click" and action not in clicked:
                    clicked.append(action)
                if event["event_type"] == "scroll" and event["event_name"] != "scroll.final" and action not in scroll_milestones:
                    scroll_milestones.append(action)

            sessions_sheet.append([
                row["attempt_uuid"] or row["session_id"],
                row["session_id"],
                source,
                utm.get("utm_source"),
                utm.get("utm_medium"),
                utm.get("utm_campaign"),
                utm.get("utm_content"),
                utm.get("utm_term"),
                row["referrer"],
                row["landing_path"],
                row["first_seen_at"],
                row["last_seen_at"],
                row["visible_scroll_percent"] or 0,
                row["click_count"] or 0,
                row["event_count"] or 0,
                ", ".join(clicked),
                ", ".join(scroll_milestones),
                self.registration_label(attempt_status, attempt_id, attempt_step),
                attempt_id,
                attempt_step,
                f'{row["viewport_width"] or ""}x{row["viewport_height"] or ""}',
                row["page_height"],
                row["user_agent"],
            ])

        setup_sheet(sessions_sheet)

        events_sheet = workbook.create_sheet("Действия")
        events_sheet.append([
            "Время",
            "UUID заявки",
            "Сессия",
            "Источник",
            "UTM source",
            "UTM medium",
            "UTM campaign",
            "Тип",
            "Действие",
            "Событие",
            "Текст",
            "Href",
            "Страница",
            "Scroll %",
            "ID заявки",
            "Регистрация",
        ])

        for row in behavior_events:
            utm = self.parse_json_object(row["session_utm_json"])
            source = self.normalized_source({"utm": utm, "referrer": row["session_referrer"]})
            attempt_id = row["linked_attempt_id"]
            attempt_status = row["attempt_status"] or ("in_progress" if attempt_id else "not_started")
            events_sheet.append([
                row["created_at"],
                row["attempt_uuid"] or row["session_id"],
                row["session_id"],
                source,
                utm.get("utm_source"),
                utm.get("utm_medium"),
                utm.get("utm_campaign"),
                row["event_type"],
                self.behavior_action_label(row),
                row["event_name"],
                row["target_text"],
                row["target_href"],
                row["page_path"] or row["landing_path"],
                row["scroll_percent"],
                attempt_id,
                self.registration_label(attempt_status, attempt_id, row["attempt_step"] or 0),
            ])

        setup_sheet(events_sheet)

        sources: dict[str, dict] = {}
        for row in behavior_sessions:
            utm = self.parse_json_object(row["utm_json"])
            source = self.normalized_source({"utm": utm, "referrer": row["referrer"]})
            bucket = sources.setdefault(
                source,
                {
                    "sessions": 0,
                    "clicks": 0,
                    "submitted": 0,
                    "scroll_total": 0.0,
                    "scroll25": 0,
                    "scroll50": 0,
                    "scroll75": 0,
                    "scroll90": 0,
                },
            )
            scroll = float(row["visible_scroll_percent"] or 0)
            attempt_id = row["attempt_id"] or row["registration_attempt_id"] or row["event_attempt_id"]
            attempt_status = row["attempt_status"] or ("in_progress" if attempt_id else "not_started")
            bucket["sessions"] += 1
            bucket["clicks"] += int(row["click_count"] or 0)
            bucket["submitted"] += 1 if attempt_status == "submitted" else 0
            bucket["scroll_total"] += scroll
            bucket["scroll25"] += 1 if scroll >= 25 else 0
            bucket["scroll50"] += 1 if scroll >= 50 else 0
            bucket["scroll75"] += 1 if scroll >= 75 else 0
            bucket["scroll90"] += 1 if scroll >= 90 else 0

        sources_sheet = workbook.create_sheet("Источники поведения")
        sources_sheet.append([
            "Источник",
            "Посетители",
            "Клики",
            "Заявки отправлены",
            "Средний скролл %",
            "Дошли 25%",
            "Дошли 50%",
            "Дошли 75%",
            "Дошли 90%",
        ])
        for source, values in sorted(sources.items()):
            sessions_count = values["sessions"] or 1
            sources_sheet.append([
                source,
                values["sessions"],
                values["clicks"],
                values["submitted"],
                round(values["scroll_total"] / sessions_count, 1),
                values["scroll25"],
                values["scroll50"],
                values["scroll75"],
                values["scroll90"],
            ])

        setup_sheet(sources_sheet)

    def handle_delete_attempt(self, path: str) -> None:
        attempt_id = self.extract_admin_attempt_id(path)
        if attempt_id is None:
            self.send_json({"error": "Invalid attempt id"}, HTTPStatus.BAD_REQUEST)
            return

        with sqlite3.connect(DB_PATH) as connection:
            cursor = connection.execute("DELETE FROM registration_attempts WHERE id = ?", (attempt_id,))
            connection.commit()

        self.send_json({"deleted": cursor.rowcount, "ids": [attempt_id]})

    def handle_delete_attempts(self) -> None:
        body = self.read_json()
        raw_ids = body.get("ids") if isinstance(body.get("ids"), list) else []
        ids = []

        for raw_id in raw_ids:
            try:
                attempt_id = int(raw_id)
            except (TypeError, ValueError):
                continue

            if attempt_id > 0 and attempt_id not in ids:
                ids.append(attempt_id)

        if not ids:
            self.send_json({"error": "No ids provided"}, HTTPStatus.BAD_REQUEST)
            return

        placeholders = ",".join("?" for _ in ids)
        with sqlite3.connect(DB_PATH) as connection:
            cursor = connection.execute(
                f"DELETE FROM registration_attempts WHERE id IN ({placeholders})",
                ids,
            )
            connection.commit()

        self.send_json({"deleted": cursor.rowcount, "ids": ids})

    @staticmethod
    def extract_attempt_id(path: str) -> int | None:
        parts = [part for part in path.split("/") if part]
        if len(parts) < 3:
            return None

        try:
            return int(parts[2])
        except ValueError:
            return None

    @staticmethod
    def extract_admin_attempt_id(path: str) -> int | None:
        parts = [part for part in path.split("/") if part]
        if len(parts) < 4:
            return None

        try:
            return int(parts[3])
        except ValueError:
            return None

    @staticmethod
    def normalized_source(item: dict) -> str:
        utm = item.get("utm") if isinstance(item.get("utm"), dict) else {}
        referrer = str(item.get("referrer") or "").lower()
        utm_source = PlankaHandler.normalize_source_token(utm.get("utm_source"))
        values = [utm.get("utm_region"), utm_source, utm.get("utm_medium"), utm.get("utm_campaign")]
        source_text = " ".join(str(value).lower() for value in values if value)
        combined = f"{source_text} {referrer}"

        if "yandex" in combined or utm.get("yclid"):
            return "yandex"

        if "telegram" in combined or "t.me" in combined or "tg" in combined:
            return "telegram"

        if "google" in combined or utm.get("gclid"):
            return "google"

        if "vk" in combined:
            return "vk"

        if utm_source:
            return utm_source

        referrer_source = PlankaHandler.referrer_source(referrer)
        if referrer_source:
            return referrer_source

        if any(str(value or "").strip() for value in utm.values()):
            return "не указан"

        return "прямой"

    @staticmethod
    def normalize_source_token(value) -> str:
        if value is None:
            return ""

        source = str(value).strip().lower()
        if not source:
            return ""

        aliases = {
            "ya": "yandex",
            "яндекс": "yandex",
            "tg": "telegram",
            "телеграм": "telegram",
            "whats app": "whatsapp",
            "whats-app": "whatsapp",
            "whatapp": "whatsapp",
            "watsapp": "whatsapp",
            "wa": "whatsapp",
        }
        return aliases.get(source, source.replace(" ", "_"))

    @staticmethod
    def referrer_source(referrer: str) -> str:
        if not referrer:
            return ""

        host = urlparse(referrer).netloc.lower()
        if not host:
            return ""

        if host.startswith("www."):
            host = host[4:]

        if host.startswith("127.0.0.1") or host.startswith("localhost") or "plankahub." in host:
            return ""

        host_without_port = host.split(":", 1)[0]
        host_parts = host_without_port.split(".")
        if len(host_parts) == 4 and all(part.isdigit() for part in host_parts):
            return ""

        return host

    @staticmethod
    def analytics_sources(attempts: list[dict]) -> list[str]:
        source_set = {PlankaHandler.normalized_source(item) for item in attempts}
        preferred = ["yandex", "telegram", "google", "vk", "whatsapp", "прямой", "не указан"]
        sources = [source for source in preferred if source in source_set]
        sources.extend(sorted(source for source in source_set if source not in preferred))
        return sources or ["не указан"]

    @staticmethod
    def add_analytics_sheet_v2(
        workbook,
        attempts: list[dict],
        column_letters: dict[str, str],
        Alignment,
        BarChart,
        Border,
        Font,
        PatternFill,
        PieChart,
        Reference,
        Side,
    ) -> None:
        sheet = workbook.create_sheet("Аналитика")
        chart_data = workbook.create_sheet("_Данные диаграмм")
        chart_data.sheet_state = "hidden"

        total_attempts = len(attempts)
        sources = PlankaHandler.analytics_sources(attempts)
        source_values = {
            source: sum(1 for item in attempts if PlankaHandler.normalized_source(item) == source)
            for source in sources
        }
        status_values = {
            source: {
                "В процессе": sum(
                    1
                    for item in attempts
                    if PlankaHandler.normalized_source(item) == source and item["status"] != "submitted"
                ),
                "Отправлено": sum(
                    1
                    for item in attempts
                    if PlankaHandler.normalized_source(item) == source and item["status"] == "submitted"
                ),
            }
            for source in sources
        }

        funnel_values = [
            ("Открыли форму (шаг 0+)", total_attempts),
            ("Дошли до шага 1", sum(1 for item in attempts if int(item["currentStep"] or 0) >= 1)),
            ("Дошли до шага 2", sum(1 for item in attempts if int(item["currentStep"] or 0) >= 2)),
            ("Дошли до шага 3", sum(1 for item in attempts if int(item["currentStep"] or 0) >= 3)),
            ("Отправили заявку", sum(1 for item in attempts if item["status"] == "submitted")),
        ]
        dropoff_values = [
            ("На шаге 0 → 1", funnel_values[0][1] - funnel_values[1][1], funnel_values[0][1]),
            ("На шаге 1 → 2", funnel_values[1][1] - funnel_values[2][1], funnel_values[1][1]),
            ("На шаге 2 → 3", funnel_values[2][1] - funnel_values[3][1], funnel_values[2][1]),
            ("При отправке", funnel_values[3][1] - funnel_values[4][1], funnel_values[3][1]),
        ]
        field_values = [
            ("Имя", "first-name"),
            ("Фамилия", "last-name"),
            ("Роль", "role"),
            ("Стадия проекта", "stage"),
            ("Потребности", "needs"),
            ("Сумма запроса", "request"),
            ("Email", "email"),
            ("Контакт", "contact"),
            ("Согласие (privacy)", "privacy"),
        ]

        utm_value_counts: dict[tuple[str, str], int] = {}
        utm_key_counts: dict[str, int] = {}
        for item in attempts:
            utm = item.get("utm") if isinstance(item.get("utm"), dict) else {}
            for key, value in utm.items():
                if value is None or value == "":
                    continue

                value_text = str(value).strip()
                if not value_text:
                    continue

                key_text = str(key)
                utm_key_counts[key_text] = utm_key_counts.get(key_text, 0) + 1
                counter_key = (key_text, value_text)
                utm_value_counts[counter_key] = utm_value_counts.get(counter_key, 0) + 1

        table_ranges: list[tuple[int, int, int]] = []
        section_rows: list[int] = []
        header_rows: list[int] = []
        percent_cells: list[tuple[int, int]] = []

        def add_section(title: str, headers: list[str], rows: list[list], percent_columns: tuple[int, ...] = ()) -> tuple[int, int, int]:
            sheet.append([])
            section_row = sheet.max_row + 1
            sheet.append([title])
            header_row = sheet.max_row + 1
            sheet.append(headers)
            first_data_row = sheet.max_row + 1
            for row in rows:
                sheet.append(row)
            last_data_row = sheet.max_row

            section_rows.append(section_row)
            header_rows.append(header_row)
            table_ranges.append((header_row, last_data_row, len(headers)))
            for row_index in range(first_data_row, last_data_row + 1):
                for column_index in percent_columns:
                    percent_cells.append((row_index, column_index))

            return header_row, first_data_row, last_data_row

        sheet.append(["Аналитика по попыткам регистрации"])
        sheet.append(["Всего попыток:", total_attempts])

        source_rows = [
            [source, source_values[source], source_values[source] / total_attempts if total_attempts else 0]
            for source in sources
        ]
        source_rows.append(["Итого", sum(source_values.values()), 1 if total_attempts else 0])
        source_header_row, source_first_row, source_last_row = add_section(
            "1. Источники трафика",
            ["Источник", "Кол-во", "Доля"],
            source_rows,
            (3,),
        )

        status_rows = [
            [source, status_values[source]["В процессе"], status_values[source]["Отправлено"]]
            for source in sources
        ]
        source_status_header_row, _source_status_first_row, source_status_last_row = add_section(
            "2. Источники × Статус заявки",
            ["Источник", "В процессе", "Отправлено"],
            status_rows,
        )

        funnel_rows = [
            [title, value, value / total_attempts if total_attempts else 0]
            for title, value in funnel_values
        ]
        _funnel_header_row, _funnel_first_row, _funnel_last_row = add_section(
            "3. Воронка по шагам",
            ["Этап", "Пользователей", "Конверсия от старта"],
            funnel_rows,
            (3,),
        )

        dropoff_rows = [
            [title, value, value / base if base else 0]
            for title, value, base in dropoff_values
        ]
        _dropoff_header_row, _dropoff_first_row, _dropoff_last_row = add_section(
            "4. Отвалы на каждом шаге",
            ["Переход", "Отвалилось", "% отвала"],
            dropoff_rows,
            (3,),
        )

        field_rows = []
        for title, key in field_values:
            filled = sum(1 for item in attempts if PlankaHandler.field_is_filled(item, key))
            field_rows.append([title, filled, filled / total_attempts if total_attempts else 0])
        _field_header_row, _field_first_row, _field_last_row = add_section(
            "5. Заполняемость полей",
            ["Поле", "Заполнили", "% заполнения"],
            field_rows,
            (3,),
        )

        utm_key_rows = [
            [key, count, count / total_attempts if total_attempts else 0]
            for key, count in sorted(utm_key_counts.items())
        ] or [["-", 0, 0]]
        _utm_key_header_row, _utm_key_first_row, _utm_key_last_row = add_section(
            "6. Заполняемость UTM-параметров",
            ["Параметр", "Заполнено", "% от попыток"],
            utm_key_rows,
            (3,),
        )

        utm_value_rows = [
            [key, value, count, count / total_attempts if total_attempts else 0]
            for (key, value), count in sorted(utm_value_counts.items(), key=lambda item: (item[0][0], item[0][1]))
        ] or [["-", "нет данных", 0, 0]]
        _utm_value_header_row, _utm_value_first_row, _utm_value_last_row = add_section(
            "7. Значения UTM-параметров",
            ["Параметр", "Значение", "Кол-во", "Доля"],
            utm_value_rows,
            (4,),
        )

        chart_data.append(["Источник", "Кол-во", "Доля"])
        for source in sources:
            count = source_values[source]
            chart_data.append([source, count, count / total_attempts if total_attempts else 0])

        chart_data["E1"] = "Источник"
        chart_data["F1"] = "В процессе"
        chart_data["G1"] = "Отправлено"
        for index, source in enumerate(sources, start=2):
            chart_data.cell(index, 5, source)
            chart_data.cell(index, 6, status_values[source]["В процессе"])
            chart_data.cell(index, 7, status_values[source]["Отправлено"])

        chart_data["I1"] = "Этап"
        chart_data["J1"] = "Пользователей"
        for index, (title, value) in enumerate(funnel_values, start=2):
            chart_data.cell(index, 9, title)
            chart_data.cell(index, 10, value)

        chart_data["L1"] = "Переход"
        chart_data["M1"] = "Отвалилось"
        for index, (title, value, _base) in enumerate(dropoff_values, start=2):
            chart_data.cell(index, 12, title)
            chart_data.cell(index, 13, value)

        chart_data["O1"] = "Поле"
        chart_data["P1"] = "Заполнили"
        for index, (title, key) in enumerate(field_values, start=2):
            chart_data.cell(index, 15, title)
            chart_data.cell(index, 16, sum(1 for item in attempts if PlankaHandler.field_is_filled(item, key)))

        blue = "1F4E79"
        light_blue = "D9EAF7"
        border_side = Side(style="thin", color="BFBFBF")
        table_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        header_fill = PatternFill("solid", fgColor=blue)
        total_fill = PatternFill("solid", fgColor=light_blue)

        sheet["A1"].font = Font(name="Arial", size=14, bold=True, color=blue)
        sheet["A2"].font = Font(name="Arial", bold=True)
        sheet["B2"].font = Font(name="Arial", size=11, bold=True, color=blue)

        for row_index in section_rows:
            sheet.cell(row_index, 1).font = Font(name="Arial", size=11, bold=True)

        for row_index in header_rows:
            max_column = sheet.max_column
            for column_index in range(1, max_column + 1):
                cell = sheet.cell(row_index, column_index)
                if cell.value is None:
                    continue
                cell.fill = header_fill
                cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = table_border

        for first_row, last_row, max_column in table_ranges:
            for row in sheet.iter_rows(min_row=first_row, max_row=last_row, min_col=1, max_col=max_column):
                for cell in row:
                    cell.border = table_border
                    cell.font = Font(name="Arial", size=10)
                    cell.alignment = Alignment(
                        horizontal="center" if cell.column > 1 else "left",
                        vertical="center",
                        wrap_text=True,
                    )

        for cell in sheet[source_last_row]:
            cell.fill = total_fill
            cell.font = Font(name="Arial", size=10, bold=True)

        for row_index, column_index in percent_cells:
            sheet.cell(row_index, column_index).number_format = "0.0%"

        sheet.column_dimensions["A"].width = 30
        sheet.column_dimensions["B"].width = 22
        sheet.column_dimensions["C"].width = 18
        sheet.column_dimensions["D"].width = 18
        sheet.column_dimensions["E"].width = 24
        sheet.column_dimensions["F"].width = 16
        sheet.column_dimensions["G"].width = 16
        sheet.freeze_panes = "A5"

        chart_source_rows = len(sources) + 1
        pie = PieChart()
        pie.title = "Источники трафика"
        pie.add_data(Reference(chart_data, min_col=2, min_row=1, max_row=chart_source_rows), titles_from_data=True)
        pie.set_categories(Reference(chart_data, min_col=1, min_row=2, max_row=chart_source_rows))
        pie.legend.position = "r"
        pie.height = 9
        pie.width = 14
        sheet.add_chart(pie, "E4")

        status_chart = BarChart()
        status_chart.title = "Источники × Статус заявки"
        status_chart.y_axis.title = "Кол-во"
        status_chart.x_axis.title = "Источник"
        status_chart.add_data(Reference(chart_data, min_col=6, max_col=7, min_row=1, max_row=chart_source_rows), titles_from_data=True)
        status_chart.set_categories(Reference(chart_data, min_col=5, min_row=2, max_row=chart_source_rows))
        status_chart.legend.position = "r"
        status_chart.height = 8.5
        status_chart.width = 14
        sheet.add_chart(status_chart, "E23")

        funnel_chart = BarChart()
        funnel_chart.type = "bar"
        funnel_chart.title = "Воронка по шагам"
        funnel_chart.y_axis.title = "Показатели"
        funnel_chart.x_axis.title = "Пользователей"
        funnel_chart.add_data(Reference(chart_data, min_col=10, min_row=1, max_row=6), titles_from_data=True)
        funnel_chart.set_categories(Reference(chart_data, min_col=9, min_row=2, max_row=6))
        funnel_chart.legend = None
        funnel_chart.height = 9
        funnel_chart.width = 14
        sheet.add_chart(funnel_chart, "K4")

        dropoff_chart = BarChart()
        dropoff_chart.title = "Отвалы на каждом шаге"
        dropoff_chart.y_axis.title = "Отвалилось"
        dropoff_chart.add_data(Reference(chart_data, min_col=13, min_row=1, max_row=5), titles_from_data=True)
        dropoff_chart.set_categories(Reference(chart_data, min_col=12, min_row=2, max_row=5))
        dropoff_chart.legend = None
        dropoff_chart.height = 8.5
        dropoff_chart.width = 14
        sheet.add_chart(dropoff_chart, "K23")

        fields_chart = BarChart()
        fields_chart.type = "bar"
        fields_chart.title = "Заполняемость полей"
        fields_chart.x_axis.title = "Кол-во"
        fields_chart.add_data(Reference(chart_data, min_col=16, min_row=1, max_row=10), titles_from_data=True)
        fields_chart.set_categories(Reference(chart_data, min_col=15, min_row=2, max_row=10))
        fields_chart.legend = None
        fields_chart.height = 9.5
        fields_chart.width = 14
        sheet.add_chart(fields_chart, "E42")

    @staticmethod
    def add_analytics_sheet(
        workbook,
        attempts: list[dict],
        column_letters: dict[str, str],
        Alignment,
        BarChart,
        Border,
        Font,
        PatternFill,
        PieChart,
        Reference,
        Side,
    ) -> None:
        sheet = workbook.create_sheet("Аналитика")
        chart_data = workbook.create_sheet("_Данные диаграмм")
        chart_data.sheet_state = "hidden"
        attempts_sheet = "'Попытки регистрации'"
        status_col = column_letters["status"]
        source_col = column_letters["source_norm"]
        step_col = column_letters["step_num"]

        def range_ref(column_key: str) -> str:
            column = column_letters[column_key]
            return f"{attempts_sheet}!${column}$2:${column}$5000"

        def source_count(source: str) -> str:
            return f'=COUNTIF({range_ref("source_norm")},"{source}")'

        def source_status_count(source: str, status: str) -> str:
            return (
                f'=COUNTIFS({attempts_sheet}!${source_col}$2:${source_col}$5000,"{source}",'
                f'{attempts_sheet}!${status_col}$2:${status_col}$5000,"{status}")'
            )

        def step_count(step: int) -> str:
            return f'=COUNTIF({attempts_sheet}!${step_col}$2:${step_col}$5000,">={step}")'

        rows = [
            ["Аналитика по попыткам регистрации"],
            ["Всего попыток:", f'=COUNTA({range_ref("id")})'],
            [],
            ["1. Источники трафика"],
            ["Источник", "Кол-во", "Доля"],
        ]

        sources = ["yandex", "telegram", "google", "vk", "прямой", "не указан"]
        for row_index, source in enumerate(sources, start=6):
            rows.append([source, source_count(source), f"=IFERROR(B{row_index}/SUM($B$6:$B$11),0)"])

        rows.extend([
            ["Итого", "=SUM(B6:B11)", "=SUM(C6:C11)"],
            [],
            [],
            ["2. Источники × Статус заявки"],
            ["Источник", "В процессе", "Отправлено"],
        ])

        for source in sources:
            rows.append([
                source,
                source_status_count(source, "В процессе"),
                source_status_count(source, "Отправлено"),
            ])

        rows.extend([
            [],
            [],
            ["3. Воронка по шагам"],
            ["Этап", "Пользователей", "Конверсия от старта"],
            ["Открыли форму (шаг 0+)", f'=COUNTA({range_ref("id")})', "=IFERROR(B27/$B$27,0)"],
            ["Дошли до шага 1", step_count(1), "=IFERROR(B28/$B$27,0)"],
            ["Дошли до шага 2", step_count(2), "=IFERROR(B29/$B$27,0)"],
            ["Дошли до шага 3", step_count(3), "=IFERROR(B30/$B$27,0)"],
            ["Отправили заявку", f'=COUNTIF({range_ref("status")},"Отправлено")', "=IFERROR(B31/$B$27,0)"],
            [],
            [],
            ["4. Отвалы на каждом шаге"],
            ["Переход", "Отвалилось", "% отвала"],
            ["На шаге 0 → 1", "=B27-B28", "=IFERROR(B36/B27,0)"],
            ["На шаге 1 → 2", "=B28-B29", "=IFERROR(B37/B28,0)"],
            ["На шаге 2 → 3", "=B29-B30", "=IFERROR(B38/B29,0)"],
            ["При отправке", "=B30-B31", "=IFERROR(B39/B30,0)"],
            [],
            [],
            ["5. Заполняемость полей"],
            ["Поле", "Заполнили", "% заполнения"],
        ])

        fill_rows = [
            ("Имя", "first-name", 44),
            ("Фамилия", "last-name", 45),
            ("Роль", "role", 46),
            ("Стадия проекта", "stage", 47),
            ("Потребности", "needs", 48),
            ("Сумма запроса", "request", 49),
            ("Email", "email", 50),
            ("Контакт", "contact", 51),
            ("Согласие (privacy)", "privacy", 52),
        ]
        for title, column_key, row_index in fill_rows:
            rows.append([title, f'=COUNTA({range_ref(column_key)})', f"=IFERROR(B{row_index}/$B$2,0)"])

        for row in rows:
            sheet.append(row)

        sources = ["yandex", "telegram", "google", "vk", "прямой", "не указан"]
        source_values = {
            source: sum(1 for item in attempts if PlankaHandler.normalized_source(item) == source)
            for source in sources
        }
        total_attempts = len(attempts)
        status_values = {
            source: {
                "В процессе": sum(
                    1
                    for item in attempts
                    if PlankaHandler.normalized_source(item) == source and item["status"] != "submitted"
                ),
                "Отправлено": sum(
                    1
                    for item in attempts
                    if PlankaHandler.normalized_source(item) == source and item["status"] == "submitted"
                ),
            }
            for source in sources
        }
        funnel_values = [
            ("Открыли форму (шаг 0+)", total_attempts),
            ("Дошли до шага 1", sum(1 for item in attempts if int(item["currentStep"] or 0) >= 1)),
            ("Дошли до шага 2", sum(1 for item in attempts if int(item["currentStep"] or 0) >= 2)),
            ("Дошли до шага 3", sum(1 for item in attempts if int(item["currentStep"] or 0) >= 3)),
            ("Отправили заявку", sum(1 for item in attempts if item["status"] == "submitted")),
        ]
        dropoff_values = [
            ("На шаге 0 → 1", funnel_values[0][1] - funnel_values[1][1]),
            ("На шаге 1 → 2", funnel_values[1][1] - funnel_values[2][1]),
            ("На шаге 2 → 3", funnel_values[2][1] - funnel_values[3][1]),
            ("При отправке", funnel_values[3][1] - funnel_values[4][1]),
        ]
        field_values = [
            ("Имя", "first-name"),
            ("Фамилия", "last-name"),
            ("Роль", "role"),
            ("Стадия проекта", "stage"),
            ("Потребности", "needs"),
            ("Сумма запроса", "request"),
            ("Email", "email"),
            ("Контакт", "contact"),
            ("Согласие (privacy)", "privacy"),
        ]

        chart_data.append(["Источник", "Кол-во", "Доля"])
        for source in sources:
            count = source_values[source]
            chart_data.append([source, count, count / total_attempts if total_attempts else 0])

        chart_data["E1"] = "Источник"
        chart_data["F1"] = "В процессе"
        chart_data["G1"] = "Отправлено"
        for index, source in enumerate(sources, start=2):
            chart_data.cell(index, 5, source)
            chart_data.cell(index, 6, status_values[source]["В процессе"])
            chart_data.cell(index, 7, status_values[source]["Отправлено"])

        chart_data["I1"] = "Этап"
        chart_data["J1"] = "Пользователей"
        for index, (title, value) in enumerate(funnel_values, start=2):
            chart_data.cell(index, 9, title)
            chart_data.cell(index, 10, value)

        chart_data["L1"] = "Переход"
        chart_data["M1"] = "Отвалилось"
        for index, (title, value) in enumerate(dropoff_values, start=2):
            chart_data.cell(index, 12, title)
            chart_data.cell(index, 13, value)

        chart_data["O1"] = "Поле"
        chart_data["P1"] = "Заполнили"
        for index, (title, key) in enumerate(field_values, start=2):
            chart_data.cell(index, 15, title)
            chart_data.cell(index, 16, sum(1 for item in attempts if PlankaHandler.field_is_filled(item, key)))

        blue = "1F4E79"
        light_blue = "D9EAF7"
        border_side = Side(style="thin", color="BFBFBF")
        table_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        header_fill = PatternFill("solid", fgColor=blue)
        total_fill = PatternFill("solid", fgColor=light_blue)

        sheet["A1"].font = Font(name="Arial", size=14, bold=True, color=blue)
        sheet["A2"].font = Font(name="Arial", bold=True)
        sheet["B2"].font = Font(name="Arial", size=11, bold=True, color=blue)

        section_rows = [4, 15, 25, 34, 42]
        header_rows = [5, 16, 26, 35, 43]
        for row_index in section_rows:
            sheet.cell(row_index, 1).font = Font(name="Arial", size=11, bold=True)

        for row_index in header_rows:
            for column_index in range(1, 4):
                cell = sheet.cell(row_index, column_index)
                cell.fill = header_fill
                cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = table_border

        table_ranges = [(6, 12), (17, 22), (27, 31), (36, 39), (44, 52)]
        for first_row, last_row in table_ranges:
            for row in sheet.iter_rows(min_row=first_row, max_row=last_row, min_col=1, max_col=3):
                for cell in row:
                    cell.border = table_border
                    cell.font = Font(name="Arial", size=10)
                    cell.alignment = Alignment(
                        horizontal="center" if cell.column > 1 else "left",
                        vertical="center",
                        wrap_text=True,
                    )

        for cell in sheet[12]:
            cell.fill = total_fill
            cell.font = Font(name="Arial", size=10, bold=True)

        for row_index in [6, 7, 8, 9, 10, 11, 12, 27, 28, 29, 30, 31, 36, 37, 38, 39, 44, 45, 46, 47, 48, 49, 50, 51, 52]:
            sheet.cell(row_index, 3).number_format = "0.0%"

        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 15
        sheet.column_dimensions["C"].width = 20
        sheet.column_dimensions["E"].width = 24
        sheet.column_dimensions["F"].width = 16
        sheet.column_dimensions["G"].width = 16
        sheet.freeze_panes = "A5"

        pie = PieChart()
        pie.title = "Источники трафика"
        pie.add_data(Reference(chart_data, min_col=2, min_row=1, max_row=7), titles_from_data=True)
        pie.set_categories(Reference(chart_data, min_col=1, min_row=2, max_row=7))
        pie.legend.position = "r"
        pie.height = 9
        pie.width = 14
        sheet.add_chart(pie, "E4")

        status_chart = BarChart()
        status_chart.title = "Источники × Статус заявки"
        status_chart.y_axis.title = "Кол-во"
        status_chart.x_axis.title = "Источник"
        status_chart.add_data(Reference(chart_data, min_col=6, max_col=7, min_row=1, max_row=7), titles_from_data=True)
        status_chart.set_categories(Reference(chart_data, min_col=5, min_row=2, max_row=7))
        status_chart.legend.position = "r"
        status_chart.height = 8.5
        status_chart.width = 14
        sheet.add_chart(status_chart, "E23")

        funnel_chart = BarChart()
        funnel_chart.type = "bar"
        funnel_chart.title = "Воронка по шагам"
        funnel_chart.y_axis.title = "Показатели"
        funnel_chart.x_axis.title = "Пользователей"
        funnel_chart.add_data(Reference(chart_data, min_col=10, min_row=1, max_row=6), titles_from_data=True)
        funnel_chart.set_categories(Reference(chart_data, min_col=9, min_row=2, max_row=6))
        funnel_chart.legend = None
        funnel_chart.height = 9
        funnel_chart.width = 14
        sheet.add_chart(funnel_chart, "K4")

        dropoff_chart = BarChart()
        dropoff_chart.title = "Отвалы на каждом шаге"
        dropoff_chart.y_axis.title = "Отвалилось"
        dropoff_chart.add_data(Reference(chart_data, min_col=13, min_row=1, max_row=5), titles_from_data=True)
        dropoff_chart.set_categories(Reference(chart_data, min_col=12, min_row=2, max_row=5))
        dropoff_chart.legend = None
        dropoff_chart.height = 8.5
        dropoff_chart.width = 14
        sheet.add_chart(dropoff_chart, "K23")

        fields_chart = BarChart()
        fields_chart.type = "bar"
        fields_chart.title = "Заполняемость полей"
        fields_chart.x_axis.title = "Кол-во"
        fields_chart.add_data(Reference(chart_data, min_col=16, min_row=1, max_row=10), titles_from_data=True)
        fields_chart.set_categories(Reference(chart_data, min_col=15, min_row=2, max_row=10))
        fields_chart.legend = None
        fields_chart.height = 9.5
        fields_chart.width = 14
        sheet.add_chart(fields_chart, "E42")

    @staticmethod
    def field_is_filled(item: dict, key: str) -> bool:
        fields = item.get("fields") if isinstance(item.get("fields"), dict) else {}
        value = fields.get(key)

        if value is None or value == "":
            return False

        if isinstance(value, list):
            return len(value) > 0

        return True

    @staticmethod
    def excel_value(value) -> str | int | float:
        if value is None or value == "":
            return None

        if isinstance(value, bool):
            return "да" if value else "нет"

        if isinstance(value, (dict, list)):
            return json.dumps(value, ensure_ascii=False)

        return value


def run() -> None:
    init_db()
    host = "127.0.0.1"
    port = 4173
    server = ThreadingHTTPServer((host, port), PlankaHandler)
    print(f"Planka Hub server: http://{host}:{port}/")
    print(f"Admin panel:       http://{host}:{port}{ADMIN_PATH}")
    server.serve_forever()


if __name__ == "__main__":
    run()
